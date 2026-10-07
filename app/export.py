"""Deterministic PDF/DOCX export. No LLM involved."""
from __future__ import annotations
from io import BytesIO
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from docx import Document
from app.services.document import render_document
from app.storage.sessions import Session


def pdf_bytes(session: Session) -> bytes:
    out = BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4, rightMargin=45, leftMargin=45, topMargin=45, bottomMargin=45)
    styles = getSampleStyleSheet()
    story=[]
    for line in render_document(session.state).splitlines():
        if line.startswith("# "):
            story.append(Paragraph(line[2:], styles["Title"]))
        elif line.startswith("## "):
            story.append(Paragraph(line[3:], styles["Heading2"]))
        elif line.startswith("> "):
            story.append(Paragraph(line[2:], styles["Italic"]))
        elif line.startswith("- "):
            story.append(Paragraph("• " + line[2:], styles["BodyText"]))
        elif line:
            story.append(Paragraph(line.replace("&", "&amp;"), styles["BodyText"]))
        story.append(Spacer(1, 5))
    doc.build(story)
    return out.getvalue()


def docx_bytes(session: Session) -> bytes:
    out = BytesIO(); doc = Document()
    for line in render_document(session.state).splitlines():
        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line.startswith("> "):
            doc.add_paragraph(line[2:]).italic = True
        elif line:
            doc.add_paragraph(line)
    doc.save(out); return out.getvalue()
